"""RAW Renamer Studio — Excel Engine (openpyxl).

Чтение: A (1) — Код LM, D (4) — GTIN.
Запись: L (12) — Итого ракурсов (БЕЗ фото ШК _y), M (13) — Ракурс_25 (1/0).

Защита от блокировки Windows Excel (master §3): перед записью проверяется
флаг эксклюзивного доступа; при PermissionError — XlsxLockedError с текстом
«Закройте файл в Excel и нажмите "Повторить"». Сохранение атомарное
(tmp + os.replace).
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Optional

COL = {
    'A': 1, 'B': 2, 'C': 3, 'D': 4, 'E': 5, 'F': 6, 'G': 7, 'H': 8, 'I': 9,
    'J': 10, 'K': 11, 'L': 12, 'M': 13, 'N': 14, 'O': 15, 'P': 16, 'Q': 17,
    'R': 18, 'S': 19, 'T': 20, 'U': 21,
}

LOCK_MSG = 'Закройте файл в Excel и нажмите «Повторить»'
_NEW_ROW_FILL = None


def _get_new_row_fill():
    global _NEW_ROW_FILL
    if _NEW_ROW_FILL is None:
        from openpyxl.styles import PatternFill
        _NEW_ROW_FILL = PatternFill('solid', fgColor='FFF3C4')
    return _NEW_ROW_FILL


class XlsxLockedError(Exception):
    """Файл открыт в Microsoft Excel (эксклюзивная блокировка)."""


def _norm(v) -> Optional[str]:
    """Номер ячейки в строку: float 4.65e+12 → '4650101098817', int → '89458028'."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s or None


class ExcelWorker:
    def __init__(self, path):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f'Заявка не найдена: {self.path}')
        self._check_writable()
        from openpyxl import load_workbook
        self.wb = load_workbook(self.path)
        self.ws = self.wb.active
        self.by_lm: dict[str, int] = {}
        self.by_gtin: dict[str, int] = {}
        self._reindex()

    # ------------------------------------------------------------------ utils
    def _check_writable(self) -> None:
        try:
            with open(self.path, 'r+b'):
                pass
        except PermissionError as e:
            raise XlsxLockedError(LOCK_MSG) from e

    def _reindex(self) -> None:
        self.by_lm.clear()
        self.by_gtin.clear()
        for r in range(2, self.ws.max_row + 1):
            a = _norm(self.ws.cell(r, COL['A']).value)
            d = _norm(self.ws.cell(r, COL['D']).value)
            if a:
                self.by_lm.setdefault(a, r)
            if d and d.isdigit():
                self.by_gtin.setdefault(d, r)

    def save_atomic(self) -> None:
        tmp = self.path.with_name(self.path.name + '.tmp')
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as e:
            tmp.unlink(missing_ok=True)
            raise XlsxLockedError(LOCK_MSG) from e
        self._reindex()

    # ------------------------------------------------------------------- API
    def find_row(self, lm=None, gtin=None) -> Optional[int]:
        if lm:
            r = self.by_lm.get(_norm(lm))
            if r:
                return r
        if gtin:
            d = _norm(gtin)
            if d and d.isdigit():
                return self.by_gtin.get(d)
        return None

    def write_results(self, updates: list[dict], constants: dict) -> list[dict]:
        """updates: [{'lm','gtin','name','L','M'}].

        Возвращает журнал изменений (для Undo):
        [{'row','added','old':{'L','M'},'new':{'L','M'}}].
        """
        changed: list[dict] = []
        for u in updates:
            row = self.find_row(u.get('lm'), u.get('gtin'))
            added = False
            if row is None:
                # артикула нет в заявке (найден только через API) — дописываем строку
                row = self.ws.max_row + 1
                added = True
                if u.get('lm'):
                    self.ws.cell(row, COL['A'], str(u['lm']))
                if u.get('name'):
                    self.ws.cell(row, COL['C'], str(u['name']))
                if u.get('gtin'):
                    self.ws.cell(row, COL['D'], str(u['gtin']))
                self.ws.cell(row, COL['K'], constants.get('org', 'photo production'))
                if constants.get('photographer'):
                    self.ws.cell(row, COL['P'], constants['photographer'])
                for c in range(1, 22):
                    self.ws.cell(row, c).fill = _get_new_row_fill()
            old_l = self.ws.cell(row, COL['L']).value
            old_m = self.ws.cell(row, COL['M']).value
            self.ws.cell(row, COL['L'], int(u['L']))
            self.ws.cell(row, COL['M'], int(u['M']))
            changed.append({
                'row': row, 'added': added,
                'old': {'L': old_l, 'M': old_m},
                'new': {'L': int(u['L']), 'M': int(u['M'])},
            })
        self.save_atomic()
        return changed
